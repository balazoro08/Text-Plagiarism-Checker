import os
import io

class DocumentParser:
    """Utility to parse text from plain text, PDF, and DOCX files."""

    @staticmethod
    def extract_text_from_bytes(file_bytes: bytes, filename: str) -> str:
        ext = os.path.splitext(filename)[1].lower()

        if ext == '.pdf':
            return DocumentParser._parse_pdf(file_bytes)
        elif ext in ['.docx', '.doc']:
            return DocumentParser._parse_docx(file_bytes)
        else:
            return DocumentParser._parse_txt(file_bytes)

    @staticmethod
    def _parse_txt(file_bytes: bytes) -> str:
        for encoding in ['utf-8', 'latin-1', 'cp1252']:
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode('utf-8', errors='replace')

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        text = ""
        # Try PyMuPDF (fitz) first if available, fallback to pypdf
        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                text += page.get_text() + "\n"
            if text.strip():
                return text.strip()
        except ImportError:
            pass
        except Exception:
            pass

        # Fallback to pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
            return text.strip()
        except Exception as e:
            raise ValueError(f"Failed to parse PDF file: {str(e)}")

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            full_text.append(cell.text)
            return "\n".join(full_text).strip()
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file: {str(e)}")
